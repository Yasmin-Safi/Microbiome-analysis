exec(open('prep.py').read())
import json
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
out='/mnt/user-data/outputs/Supplementary_Tables_S6-S18_revision.xlsx'
def r(df,n=4):
    df=df.copy()
    for c in df.columns:
        if df[c].dtype.kind=='f': df[c]=df[c].round(6 if ('_p' in c or c.endswith('q') or '_q' in c or c=='p' or c=='q') else n)
    return df
q=pd.read_excel(U+'Qulaity_control_stats.xlsx',header=3).iloc[:,:7]; q=q[q.SampleName.astype(str).str.contains(r'^\d{3}\.S\d$')].set_index('SampleName')
for c in q.columns: q[c]=pd.to_numeric(q[c])
key=pd.DataFrame({'Sample':CT.index,'Patient':meta.pid.values,'Timepoint':meta.tp.values,'In_complete_cohort(n=46)':meta.complete.values})
key['Merged_read_pairs']=q['Read Count'].reindex(CT.index).values
key['Reads_in_species_table']=depth.values; key['Pct_merged_retained']=(100*key.Reads_in_species_table/key.Merged_read_pairs).round(1)
key['Observed_species']=(CT>0).sum(axis=1).values; key['SRA_run_accession']=''
qc=pd.read_csv('qc_by_tp.csv').rename(columns={'tp':'Timepoint','reads_med':'Merged reads median','reads_q1':'Merged reads Q1','reads_q3':'Merged reads Q3','len_mean':'Mean amplicon length (bp)','tab_med':'Reads in species table, median','frac_med':'Fraction retained, median','frac_q1':'Fraction retained Q1','frac_q3':'Fraction retained Q3'})
o1=json.load(open('out1.json'))['depth']
dep=pd.DataFrame([{'Test':'Kruskal-Wallis depth S1/S2/S3','p':o1['depth_KW_p']},{'Test':'Paired Wilcoxon depth S1 vs S2 (n=46)','p':o1['depth_paired_S1S2_p']},{'Test':'Paired Wilcoxon depth S2 vs S3 (n=46)','p':o1['depth_paired_S2S3_p']},{'Test':'Paired Wilcoxon depth S1 vs S3 (n=46)','p':o1['depth_paired_S1S3_p']},{'Test':'Spearman depth vs observed richness (unrarefied, all samples) rho','p':o1['spearman_depth_vs_obs_all'][1],'rho':o1['spearman_depth_vs_obs_all'][0]},{'Test':'Spearman depth vs observed richness (rarefied) rho','p':o1['spearman_depth_vs_obs_rar'][1],'rho':o1['spearman_depth_vs_obs_rar'][0]}])
ar=pd.read_csv('alpha_tests_rar.csv'); aw=pd.read_csv('alpha_tests_raw.csv'); al=pd.read_csv('alpha_lmm.csv')
pm=pd.read_csv('permanova.csv'); pw=pd.read_csv('permanova_pairwise.csv'); pdp=pd.read_csv('permdisp.csv')
dj=json.load(open('disp_by_tp.json')); dd=pd.DataFrame([{'Distance':k,'Set':s,'S1':v[s]['S1'],'S2':v[s]['S2'],'S3':v[s]['S3']} for k,v in dj.items() for s in ['all','complete']])
sp=pd.read_csv('da_species.csv'); ge=pd.read_csv('da_genus.csv')
pr=pd.concat([pd.read_csv(f'pair_{a}{b}.csv').assign(Comparison=f'{a} vs {b}') for a,b in [('S1','S2'),('S2','S3'),('S1','S3')]]).rename(columns={'sp':'Species','p':'MWU_p','l2fc':'log2FC_(later/earlier)','ma':'mean_RA_earlier','mb':'mean_RA_later','prev':'prevalence','q_all':'BH_q_all805','q_filt':'BH_q_filtered(prev>=10%)'})
at=pd.read_csv('attrition_compare.csv'); tip=pd.read_csv('tipping.csv'); core=pd.read_csv('core_sensitivity.csv')
reasons=pd.DataFrame([['S3 missing','Death',5,'5, 7, 8, 59, 63'],['S3 missing','Metastatic disease / progression',6,'11, 40, 41, 55, 60, 61'],['S3 missing','Withdrew from study',2,'43, 58'],['S3 missing','Stopped treatment and travelled abroad',1,'18'],['S3 missing','Insufficient sample',2,'22, 56'],['S2 missing','Sample not received (all three also lack S3)',3,'41, 55, 59']],columns=['Sample','Reason (sample-tracking record)','n','Case IDs'])
readme=pd.DataFrame({'Item':['Purpose','Sheets','Repeated measures','Rarefaction','Multiple testing','Compositional DA','Species filter','Attrition','Placeholders','Reproducibility'],
'Description':['Supplementary Tables S6-S18 supporting the revised manuscript (responses to reviewers).',
'S6 sequencing QC/sample key; S14 within-patient dissimilarity; S15 inverse-probability weighting and covariate-adjusted models; S16 rarefaction-depth sensitivity; S17 CLR specification sensitivity; S18 Dirichlet-CLR (ALDEx2-style) paired analysis; S7 alpha diversity (rarefied/unrarefied tests, mixed models); S8 beta diversity (PERMANOVA, pairwise, PERMDISP); S9 species differential abundance (+ pairwise); S10 attrition; S11 tipping-point; S12 genus level; S13 core sensitivity.',
'PERMANOVA used patient as blocking term with timepoint labels permuted within patients (9,999 permutations; 2,999 pairwise). Mixed models: random intercept per patient, REML.',
'7,500 reads per sample (minimum observed 7,560); alpha = mean of 20 rarefactions; beta = mean distance over 20 rarefactions; core prevalence = mean of 50 rarefactions.',
'Benjamini-Hochberg. KW_q_all = across all 805 species; *_filt = across species/genera with prevalence >=10% of samples (323 species, 203 genera).',
'CLR with pseudocount 0.5 on counts; LMM = CLR ~ timepoint + (1|patient) on all 167 samples; LMMd additionally adjusts for log10 depth. Not equivalent to ANCOM-BC.',
'Prevalence >=10% of the 167 samples.',
'Reasons from sample-tracking sheet; delta-adjusted analysis imputes missing S3 for the 13 patients with S2.',
'SRA_run_accession column in S6 sample key is blank: to be completed by the authors.',
'Python scripts in revision_analysis_scripts.zip regenerate all tables (seed-controlled; permutation p-values vary slightly by seed).']})
wp=pd.read_csv('within_patient_dist.csv'); wpp=pd.read_csv('within_patient_pairs.csv'); ipw=pd.read_csv('ipw.csv'); cov=pd.read_csv('alpha_covariate.csv'); rs=pd.read_csv('rarefaction_depth_sens.csv')
dsa=pd.read_csv('da_sensitivity.csv'); dsa.insert(0,'CLR_basis','full composition (then filtered)'); _d=json.load(open('da_summary.json'))['species']
dsa=pd.concat([pd.DataFrame([{'CLR_basis':'PRIMARY: prevalence-filtered table','pseudocount':0.5,'min_prevalence':0.10,'n_tested':323,'S2vS1':_d['LMM_S2vS1_q05'],'S3vS2':_d['LMM_S3vS2_q05'],'S3vS1':_d['LMM_S3vS1_q05']}]),dsa],ignore_index=True)
dso=pd.read_csv('da_sensitivity_overlap.csv'); ald=pd.read_csv('aldex_like.csv')
extra=[('S14a_Within_patient_dist',r(wp)),('S14b_Within_patient_pairs',r(wpp)),('S15a_IPW',r(ipw)),('S15b_Covariate_models',r(cov)),('S16_Rarefaction_depth',r(rs)),('S17a_DA_CLR_sensitivity',r(dsa)),('S17b_DA_overlap_S3vS1',r(dso)),('S18_Dirichlet_CLR',r(ald))]
sheets=[('README',readme),('S6a_QC_by_timepoint',r(qc)),('S6b_Depth_tests',r(dep)),('S6c_Sample_key',key),('S7a_Alpha_rarefied',r(ar)),('S7b_Alpha_unrarefied',r(aw)),('S7c_Alpha_mixed_models',r(al)),
('S8a_PERMANOVA',r(pm)),('S8b_PERMANOVA_pairwise',r(pw)),('S8c_PERMDISP',r(pdp)),('S8d_Centroid_dispersion',r(dd)),('S9a_DA_species',r(sp)),('S9b_Pairwise_DA',r(pr)),
('S10a_Attrition_reasons',reasons),('S10b_Attrition_compare',r(at)),('S11_Tipping_point',r(tip)),('S12_DA_genus',r(ge)),('S13_Core_sensitivity',r(core))]+extra
with pd.ExcelWriter(out,engine='openpyxl') as w:
    for n,df in sheets: df.to_excel(w,sheet_name=n,index=False)
wb=load_workbook(out)
for ws in wb:
    for c in ws[1]: c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10); c.fill=PatternFill('solid',fgColor='1F3864'); c.alignment=Alignment(wrap_text=True,vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font=Font(name='Arial',size=10)
    for i,col in enumerate(ws.columns,1):
        m=max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:60]); ws.column_dimensions[get_column_letter(i)].width=min(max(10,m+2),70 if ws.title=='README' else 38)
    ws.freeze_panes='A2'
    if ws.title=='README':
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
wb.save(out); print('ok',[ (n,len(d)) for n,d in sheets])
